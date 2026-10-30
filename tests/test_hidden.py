"""回歸測試：隱藏位置（Word 註腳／註解／修訂／欄位／屬性，PDF 書籤／附件／註解／連結／表單／XMP／頁外文字）
與行距過緊的 PDF。對應獨立審查員實測找到的問題。
執行：.venv/bin/python -W ignore tests/test_hidden.py
"""
import os, re, sys, zipfile
sys.path.insert(0, ".")
import pymupdf as fitz
from docx import Document
import docs_io, pii

os.makedirs("/tmp/hid", exist_ok=True)
NAME, ID, MAIL = "王大明", "A123456789", "wang.daming@example.com"
items = [{"text": NAME, "type": "NAME"}, {"text": ID, "type": "ID"}, {"text": MAIL, "type": "EMAIL"}, {"text": "舊名王大明", "type": "NAME"}]
FONT = "/System/Library/Fonts/Supplemental/Arial Unicode.ttf"
results = []

def report(name, ok, detail=""):
    results.append(ok); print(("PASS" if ok else "FAIL"), name, detail)

# ---------- Word
def make_docx(path):
    d = Document()
    d.add_paragraph(f"內文：{NAME}")
    d.core_properties.author = NAME; d.core_properties.title = f"{NAME}申請書"
    d.core_properties.subject = ID; d.core_properties.keywords = f"{NAME},{MAIL}"
    d.save(path)
    tmp = path + ".x"
    with zipfile.ZipFile(path) as zi, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zo:
        for it in zi.infolist():
            data = zi.read(it.filename)
            if it.filename == "word/document.xml":
                s = data.decode()
                s = s.replace("</w:body>", f'<w:p><w:del w:id="9" w:author="{NAME}" w:date="2026-01-01T00:00:00Z"><w:r><w:delText>舊名{NAME}</w:delText></w:r></w:del>'
                              f'<w:r><w:instrText>HYPERLINK "mailto:{MAIL}"</w:instrText></w:r></w:p></w:body>', 1) if "</w:sectPr></w:body>" not in s else \
                    s.replace("<w:sectPr", f'<w:p><w:del w:id="9" w:author="{NAME}" w:date="2026-01-01T00:00:00Z"><w:r><w:delText>舊名{NAME}</w:delText></w:r></w:del><w:r><w:instrText>HYPERLINK "mailto:{MAIL}"</w:instrText></w:r></w:p><w:sectPr', 1)
                data = s.encode()
            if it.filename == "docProps/app.xml":
                data = data.decode().replace("</Properties>", f"<Company>{NAME}農場</Company><Manager>{NAME}</Manager></Properties>").encode()
            zo.writestr(it, data)
        W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
        zo.writestr("word/footnotes.xml", f'<?xml version="1.0" encoding="UTF-8"?><w:footnotes {W}><w:footnote w:id="1"><w:p><w:r><w:t>註腳：{NAME}的</w:t></w:r><w:r><w:t>身分證{ID}</w:t></w:r></w:p></w:footnote></w:footnotes>')
        zo.writestr("word/comments.xml", f'<?xml version="1.0" encoding="UTF-8"?><w:comments {W}><w:comment w:id="0" w:author="{NAME}" w:initials="WD"><w:p><w:r><w:t>留言給{NAME}</w:t></w:r></w:p></w:comment></w:comments>')
        zo.writestr("docProps/thumbnail.jpeg", b"\xff\xd8fake")
    os.replace(tmp, path)

src = "/tmp/hid/in.docx"; dst = "/tmp/hid/out.docx"
make_docx(src)
mapping = pii.build_mapping(items, "label_typed")
docs_io.redact(src, dst, ".docx", mapping, [])
with zipfile.ZipFile(dst) as z:
    blob = "\n".join(z.read(n).decode("utf-8", "ignore") for n in z.namelist() if n.endswith((".xml", ".rels")))
    names = z.namelist()
for k in (NAME, ID, MAIL):
    report(f"Word 整包無殘留 {k}", k not in blob)
report("Word 縮圖已移除", not any(n.startswith("docProps/thumbnail") for n in names))
report("Word 自我檢查看得到殘留（對原檔）", bool(docs_io.residual(src, ".docx", mapping)))
report("Word 輸出可正常開啟", bool(Document(dst).paragraphs))
report("Word 輸出檢查乾淨", docs_io.residual(dst, ".docx", mapping) == [])

# ---------- PDF
def font_page(doc):
    p = doc.new_page(); p.insert_font(fontname="f", fontfile=FONT); return p

def make_pdf(path):
    d = fitz.open(); p = font_page(d)
    p.insert_text((50, 80), f"內文 {NAME} 的資料 {ID}", fontsize=12, fontname="f")
    p.insert_text((700, 100), f"頁外：{NAME} {ID}", fontsize=12, fontname="f")   # 頁面可視範圍外
    d.set_toc([[1, f"{NAME}申請書", 1]])
    d.embfile_add("附件.txt", f"{NAME} {ID} {MAIL}".encode())
    p.add_text_annot((300, 300), f"註解 {NAME} {ID}")
    p.insert_link({"kind": fitz.LINK_URI, "from": fitz.Rect(50, 400, 120, 420), "uri": f"mailto:{MAIL}?subject={NAME}"})
    w = fitz.Widget(); w.field_name = "n"; w.field_type = fitz.PDF_WIDGET_TYPE_TEXT; w.field_value = NAME; w.rect = fitz.Rect(50, 450, 200, 470)
    p.add_widget(w)
    d.set_xml_metadata(f"<x:xmpmeta xmlns:x='adobe:ns:meta/'><r>{NAME} {ID}</r></x:xmpmeta>")
    d.save(path)

src = "/tmp/hid/in.pdf"; dst = "/tmp/hid/out.pdf"
make_pdf(src)
before = docs_io.residual(src, ".pdf", mapping)
report("PDF 自我檢查看得到隱藏殘留（對原檔）", len(before) >= 8, str(before))
docs_io.redact(src, dst, ".pdf", mapping, [])
left = docs_io.residual(dst, ".pdf", mapping)
report("PDF 輸出檢查乾淨", left == [], str(left))
d = fitz.open(dst)
raw = b"".join(d.xref_stream(x) or b"" for x in range(1, d.xref_length()) if d.xref_is_stream(x)) + fitz.open(dst).tobytes()
dd = fitz.open(dst)
report("PDF 註解已清", all(NAME not in str(a.info) for a in dd[0].annots() or []))
report("PDF 連結已清", all(MAIL not in (l.get("uri") or "") for l in dd[0].get_links()))
report("PDF 表單欄位已清", all(NAME not in str(w.field_value) for w in dd[0].widgets() or []))
report("PDF 書籤已清", all(NAME not in t[1] for t in d.get_toc()))
report("PDF 附件已清", d.embfile_count() == 0)
report("PDF XMP 已清", NAME not in (d.get_xml_metadata() or ""))
report("PDF 頁外文字已清", NAME not in d[0].get_text("text", flags=docs_io._NOCLIP) and ID not in d[0].get_text("text", flags=docs_io._NOCLIP))

# ---------- 行距過緊
def make_tight(path, ratio):
    d = fitz.open(); p = font_page(d)
    for i, t in enumerate(["上一行的資料不可被刪 ABCDEFG", f"本行有 {NAME} 要處理", "下一行的資料不可被刪 HIJKLMN", "最後一行不可被刪 OPQRST"]):
        p.insert_text((50, 100 + i * 12 * ratio), t, fontsize=12, fontname="f")
    d.save(path)

for ratio in (1.5, 1.2, 1.0, 0.95):
    s, o = f"/tmp/hid/t{ratio}.pdf", f"/tmp/hid/t{ratio}_o.pdf"
    make_tight(s, ratio); w = []
    docs_io.redact(s, o, ".pdf", pii.build_mapping([items[0]], "label_typed"), w)
    ok = docs_io.integrity_ok(s, o, ".pdf", pii.build_mapping([items[0]], "label_typed"))
    report(f"行距 {ratio} 倍：沒誤刪，或至少有警告", ok or bool(w), f"integrity={ok} warnings={len(w)}")

sys.exit(0 if all(results) else 1)
