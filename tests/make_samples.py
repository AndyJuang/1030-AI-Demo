"""產生測試用的假資料文件（人名、號碼皆為虛構）。"""
import fitz
from docx import Document

TEXT = [
    "彰化縣員林市農會　補助申請案件紀錄",
    "申請人：王大明（身分證字號 A123456789），出生日期民國60年5月12日。",
    "戶籍地址：彰化縣員林市中山路二段123號；聯絡電話 0912-345-678，市話 (04)834-5678。",
    "電子郵件：wang.daming@example.com，匯款帳號 1234-5678-9012-3456。",
    "承辦人陳小華於113年5月到場勘查，農地位於雲林縣斗六市大學路100號，員林段123-4地號。",
    "王大明先生表示今年水稻種植面積約2公頃，預計補助金額新臺幣30,000元。",
]

def make_docx(path):
    d = Document()
    d.add_heading(TEXT[0], 1)
    for line in TEXT[1:3]:
        d.add_paragraph(line)
    p = d.add_paragraph()           # 故意把姓名與身分證拆成多個 run（不同格式）
    p.add_run("共同申請人：李").bold = True
    p.add_run("美玲").italic = True
    p.add_run("（身分證 B2")
    p.add_run("34567891）電話 0933-111-222")
    for line in TEXT[3:]:
        d.add_paragraph(line)
    t = d.add_table(rows=2, cols=2)
    t.style = "Table Grid"
    t.cell(0, 0).text, t.cell(0, 1).text = "姓名", "王大明"
    t.cell(1, 0).text, t.cell(1, 1).text = "電話", "0912-345-678"
    d.sections[0].header.paragraphs[0].text = "承辦人：陳小華"
    d.core_properties.author = "王大明"
    d.save(path)

def make_pdf(path):
    doc = fitz.open()
    page = doc.new_page()
    page.insert_font(fontname="cjk", fontfile="/System/Library/Fonts/Supplemental/Arial Unicode.ttf")
    y = 80
    for line in TEXT:
        page.insert_text((50, y), line, fontsize=11, fontname="cjk")
        y += 28
    # 會自動換行的長地址，測試跨行比對
    page.insert_textbox(fitz.Rect(50, y, 300, y + 80), "通訊地址：臺中市西屯區台灣大道三段九十九號十二樓之三，李美玲收。", fontsize=11, fontname="cjk")
    doc.subset_fonts()
    doc.save(path, garbage=4, deflate=True)

if __name__ == "__main__":
    make_docx("samples/範例_申請書.docx")
    make_pdf("samples/範例_申請書.pdf")
    print("ok")
